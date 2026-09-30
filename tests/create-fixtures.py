from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from docx import Document
from pptx import Presentation
from pptx.util import Inches

root=Path('tmp/conversion-tests'); root.mkdir(parents=True,exist_ok=True)
fontpath=r'C:\Windows\Fonts\arial.ttf'
pdfmetrics.registerFont(TTFont('ArialTest',fontpath))
for name,count in [('table',2),('second',1)]:
    c=canvas.Canvas(str(root/f'{name}.pdf'),pageCompression=0)
    for page in range(count):
        c.setFont('ArialTest',14); c.drawString(45,790,f'Convertory Test Page {page+1}')
        c.drawString(45,765,'Türkçe: İstanbul, çalışma, ödeme')
        for row,values in enumerate([['Product','Quantity','Total'],['Apple','12','120'],['Orange','8','80']]):
            y=710-row*30
            for col,value in enumerate(values): c.drawString(55+col*150,y,value)
        for x in (45,195,345,495): c.line(x,640,x,735)
        for y in (645,675,705,735): c.line(45,y,495,y)
        c.showPage()
    c.save()
im=Image.new('RGB',(1400,600),'white'); draw=ImageDraw.Draw(im)
font=ImageFont.truetype(fontpath,44)
for y,line in enumerate(['INVOICE 2026','Product          Quantity          Total','Apple                12                120','Orange               8                  80']): draw.text((50,50+y*110),line,font=font,fill='black')
im.save(root/'invoice.png')
wb=Workbook(); ws=wb.active; ws.title='Orders'; ws.append(['Product','Quantity','Total'])
for i in range(1,46): ws.append([f'Product {i}',i,i*10])
ws.column_dimensions['A'].width=20
ws=wb.create_sheet('Second Sheet'); ws.append(['SECOND SHEET MUST SURVIVE',999])
wb.save(root/'workbook.xlsx')
d=Document(); d.add_heading('Convertory Word Test',0); d.add_paragraph('Türkçe: İstanbul, çalışma, ödeme.')
t=d.add_table(rows=2,cols=2); t.cell(0,0).text='Product'; t.cell(0,1).text='Total'; t.cell(1,0).text='Apple'; t.cell(1,1).text='120'
d.save(root/'document.docx')
p=Presentation()
for title in ['First slide','Second slide']:
 s=p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text=title; s.placeholders[1].text='Convertory slide content'
p.save(root/'slides.pptx')
print('Fixtures ready:',root)
