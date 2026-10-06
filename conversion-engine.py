"""Real local conversion operations. No placeholder exports."""
import csv
import io
import json
import os
import shutil
import sys
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import fitz
from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
PS = r'C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe'

def powershell(script, *args):
    result = subprocess.run([PS, '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass', '-File', str(ROOT / script), *map(str,args)], capture_output=True, timeout=120)
    if result.returncode:
        raise ValueError(result.stderr.decode('utf-8', errors='replace')[-1800:])

IS_WINDOWS = os.name == 'nt'

def office_to_pdf(kind, source, output):
    if IS_WINDOWS:
        powershell('office-to-pdf.ps1','-Kind',kind,'-InputPath',source,'-OutputPath',output)
        return
    # Linux/container engine: LibreOffice headless with an isolated profile per job.
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice: raise ValueError('The Office conversion engine is not installed.')
    with tempfile.TemporaryDirectory() as work:
        result = subprocess.run([soffice,f'-env:UserInstallation=file://{work}/profile','--headless','--norestore','--convert-to','pdf','--outdir',work,str(source)],capture_output=True,timeout=150)
        produced = Path(work)/(Path(source).stem+'.pdf')
        if result.returncode or not produced.exists():
            raise ValueError('The document could not be converted. Check that it opens normally.')
        shutil.copyfile(produced,output)

def ocr_image(image_path):
    """Returns {'text','lines':[{'words':[{'text','x','y','width','height'}]}]} like image-ocr.ps1."""
    if IS_WINDOWS:
        ocr_path = Path(image_path).with_suffix('.ocr.json')
        powershell('image-ocr.ps1','-InputPath',image_path,'-OutputPath',ocr_path)
        return json.loads(ocr_path.read_text(encoding='utf-8-sig'))
    tesseract = shutil.which('tesseract')
    if not tesseract: raise ValueError('The OCR engine is not installed.')
    result = subprocess.run([tesseract,str(image_path),'stdout','-l','eng','--psm','6','tsv'],capture_output=True,timeout=120)
    if result.returncode: raise ValueError('Text recognition failed for this image.')
    reader = csv.DictReader(io.StringIO(result.stdout.decode('utf-8',errors='replace')),delimiter='\t',quoting=csv.QUOTE_NONE)
    lines = {}
    for row in reader:
        text = (row.get('text') or '').strip()
        if not text or float(row.get('conf') or -1) < 0: continue
        key = (row['block_num'],row['par_num'],row['line_num'])
        lines.setdefault(key,[]).append({'text':text,'x':float(row['left']),'y':float(row['top']),'width':float(row['width']),'height':float(row['height'])})
    out = [{'words':words} for words in lines.values()]
    return {'text':'\n'.join(' '.join(w['text'] for w in l['words']) for l in out),'lines':out}

def watermark_font_file():
    candidates = [r'C:\Windows\Fonts\arial.ttf','/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    for candidate in candidates:
        if Path(candidate).exists(): return candidate
    raise ValueError('No font available for the watermark.')

def save_sheet(wb, output):
    for ws in wb:
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='163A63')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(65,max(14,max(len(str(c.value or '')) for c in col)+2))
    wb.save(output)

def append_text_row(ws, row):
    # Text from input documents must not become spreadsheet formulas.
    ws.append([str(v) if v is not None else '' for v in row])
    for cell in ws[ws.max_row]:
        cell.data_type = 's'

def extract_columnar_text(page):
    # Reconstructs rows/columns from word boxes so tabular PDFs (invoices,
    # price lists) paste into Excel as real columns instead of one run-on line.
    words = page.get_text('words')
    if not words:
        return ''
    words = sorted(words, key=lambda w: (round(w[1], 1), w[0]))
    rows = []
    for x0, y0, x1, y1, text, *_ in words:
        row = next((r for r in rows if abs(r['y'] - y0) <= max(2, (y1 - y0) * 0.5)), None)
        if row is None:
            row = {'y': y0, 'words': []}
            rows.append(row)
        row['words'].append((x0, x1, y1 - y0, text))
    rows.sort(key=lambda r: r['y'])
    lines = []
    for row in rows:
        items = sorted(row['words'], key=lambda w: w[0])
        parts = [items[0][3]]
        prev_x1, prev_height = items[0][1], items[0][2]
        for x0, x1, height, text in items[1:]:
            gap = x0 - prev_x1
            parts.append('\t' if gap > max(prev_height, height) * 1.6 else ' ')
            parts.append(text)
            prev_x1, prev_height = x1, height
        lines.append(''.join(parts))
    return '\n'.join(lines)

def run(job):
    tool, paths, output = job['tool'], [Path(p) for p in job['paths']], Path(job['output'])
    source = paths[0]
    options = job.get('options', {})
    note = ''
    if tool in ('Word to PDF','Excel to PDF','PowerPoint to PDF'):
        kind = {'Word to PDF':'word','Excel to PDF':'excel','PowerPoint to PDF':'powerpoint'}[tool]
        office_to_pdf(kind, source, output)
    elif tool == 'Image to PDF':
        with fitz.open() as doc:
            for path in paths:
                with Image.open(path) as image:
                    image = ImageOps.exif_transpose(image).convert('RGB')
                    import io
                    stream = io.BytesIO()
                    image.save(stream, format='PNG')
                    scale = min(1, 14400 / max(image.size))
                    page = doc.new_page(width=image.width*scale,height=image.height*scale)
                    page.insert_image(page.rect,stream=stream.getvalue())
            doc.save(output,deflate=True)
    elif tool in ('Image to Text','Image to Excel'):
        normalized = output.with_suffix('.ocr.png')
        with Image.open(source) as im:
            im = ImageOps.exif_transpose(im).convert('RGB')
            im.thumbnail((2500,2500))
            im.save(normalized)
        ocr = ocr_image(normalized)
        if not ocr['text'].strip(): raise ValueError('No text was detected in the image. Try a sharper image.')
        # Windows OCR groups 'lines' by detected text block/column, not by row, so
        # widely-spaced columns (e.g. table cells) come back as separate lines in
        # column-major order. Re-cluster words into rows by y-position instead.
        words = sorted((w for line in ocr['lines'] for w in line['words']), key=lambda w: (w['y'], w['x']))
        rows=[]
        for word in words:
            if rows and word['y']-rows[-1]['y'] < max(word['height'],rows[-1]['h'])*0.6:
                row=rows[-1]; row['words'].append(word)
                row['y']=(row['y']*row['n']+word['y'])/(row['n']+1); row['n']+=1
            else:
                rows.append({'words':[word],'y':word['y'],'h':word['height'],'n':1})
        for row in rows: row['words'].sort(key=lambda w: w['x'])
        if tool == 'Image to Text':
            output.write_text('\n'.join(' '.join(w['text'] for w in row['words']) for row in rows),encoding='utf-8')
        else:
            wb=Workbook(); ws=wb.active; ws.title='OCR'
            for row in rows:
                cells=[]; cell=''; right=None
                for word in row['words']:
                    if right is not None and word['x']-right > max(24,word['height']*1.8):
                        cells.append(cell); cell=''
                    cell=(cell+' '+word['text']).strip(); right=word['x']+word['width']
                if cell: cells.append(cell)
                if cells: append_text_row(ws,cells)
            save_sheet(wb,output)
            note='OCR cells are editable. Review recognized values before use.'
    elif tool == 'Merge PDF':
        if len(paths)<2: raise ValueError('Select at least two PDFs to merge.')
        with fitz.open() as merged:
            for path in paths:
                with fitz.open(path) as doc: merged.insert_pdf(doc)
            merged.save(output,garbage=4,deflate=True)
    else:
        with fitz.open(source) as doc:
            if doc.needs_pass: raise ValueError('This PDF is password protected.')
            if tool == 'PDF to Excel':
                wb=Workbook(); wb.remove(wb.active)
                for i,page in enumerate(doc):
                    tables=page.find_tables().tables
                    if tables:
                        for j,table in enumerate(tables):
                            ws=wb.create_sheet(f'Page {i+1} Table {j+1}')
                            for row in table.extract(): append_text_row(ws,row)
                    else:
                        lines=page.get_text('text',sort=True).splitlines()
                        if lines:
                            ws=wb.create_sheet(f'Page {i+1} Text')
                            for line in lines: append_text_row(ws,[line])
                if not wb.sheetnames: raise ValueError('No selectable text or tables found. Scanned PDFs need OCR.')
                save_sheet(wb,output)
                note='Tables extracted where detected; other pages contain text rows.'
            elif tool == 'PDF to PowerPoint':
                prs=Presentation(); prs.slide_width=Inches(doc[0].rect.width/72); prs.slide_height=Inches(doc[0].rect.height/72)
                import io
                for page in doc:
                    slide=prs.slides.add_slide(prs.slide_layouts[6])
                    image=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
                    ratio=min(prs.slide_width/page.rect.width,prs.slide_height/page.rect.height)
                    width=int(page.rect.width*ratio); height=int(page.rect.height*ratio)
                    slide.shapes.add_picture(io.BytesIO(image.tobytes('png')),(prs.slide_width-width)//2,(prs.slide_height-height)//2,width,height)
                prs.save(output)
                note='Each PDF page is a slide image; text inside the slide image is not editable.'
            elif tool == 'Compress PDF':
                data=doc.tobytes(garbage=4,deflate=True,clean=True)
                original=source.read_bytes()
                output.write_bytes(data if len(data)<len(original) else original)
                note=(f'Reduced by {100*(1-len(data)/len(original)):.1f}% without lowering image quality.' if len(data)<len(original) else 'Already optimized; original quality and size preserved.')
            elif tool == 'Watermark PDF':
                label=str(options.get('watermark','CONFIDENTIAL')).strip()[:120]
                if not label: raise ValueError('Enter watermark text.')
                font=fitz.Font(fontfile=watermark_font_file())
                for page in doc:
                    size=min(36,(page.rect.width-48)/max(font.text_length(label,fontsize=1),1))
                    writer=fitz.TextWriter(page.rect)
                    writer.append((max(12,(page.rect.width-font.text_length(label,fontsize=size))/2),page.rect.height/2),label,font=font,fontsize=size)
                    writer.write_text(page,opacity=.25,color=(.25,.3,.4),overlay=True)
                doc.save(output,garbage=4,deflate=True)
            elif tool == 'PDF to Text':
                text='\n\n'.join(extract_columnar_text(p) for p in doc)
                if not text.strip(): raise ValueError('No selectable text found. This PDF needs OCR.')
                output.write_text(text,encoding='utf-8')
                note='Columns are separated with tabs so the text pastes into Excel as real columns.'
            elif tool == 'More Formats':
                if options.get('format','png') == 'txt':
                    text='\n\n'.join(p.get_text('text',sort=True) for p in doc)
                    if not text.strip(): raise ValueError('No selectable text found. This PDF needs OCR.')
                    output.write_text(text,encoding='utf-8')
                else:
                    with ZipFile(output,'w',ZIP_DEFLATED) as archive:
                        for i,page in enumerate(doc): archive.writestr(f'page-{i+1}.png',page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).tobytes('png'))
            else: raise ValueError('Unsupported conversion.')
    if not output.exists() or not output.stat().st_size: raise ValueError('The conversion produced an empty file.')
    if output.suffix == '.pdf':
        with fitz.open(output) as check:
            if check.page_count<1: raise ValueError('The output PDF has no pages.')
    if output.suffix in ('.xlsx','.pptx','.zip'):
        with ZipFile(output) as check:
            if check.testzip(): raise ValueError('The output archive is damaged.')
    return {'note':note,'bytes':output.stat().st_size}

if __name__=='__main__':
    try: print(json.dumps(run(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8')))))
    except Exception as error:
        print(json.dumps({'error':str(error)}))
        sys.exit(1)
