"""Local editable PDF-to-DOCX conversion helper for the Convertory preview."""
from __future__ import annotations

import argparse
import copy
import sys
import tempfile
from pathlib import Path

from pdf2docx import Converter
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn

PDF_HEADER = b"%PDF-"
PDF_EOF = b"%%EOF"


def unwrap_signed_pdf(raw_bytes: bytes) -> bytes:
    """Extract the raw PDF from a CMS/PKCS#7-enveloped signed document.

    Turkish e-Devlet/GIB downloads often wrap the actual PDF in a CMS SignedData
    envelope, so the file's first bytes are ASN.1 DER rather than "%PDF-". The
    real PDF is still embedded verbatim; find its header and trim to its last
    "%%EOF" marker to recover a directly convertible PDF.
    """
    header_at = raw_bytes.find(PDF_HEADER)
    if header_at <= 0:
        return raw_bytes
    eof_at = raw_bytes.rfind(PDF_EOF)
    if eof_at == -1 or eof_at < header_at:
        return raw_bytes
    return raw_bytes[header_at:eof_at + len(PDF_EOF)]


def remove_overflowing_page_labels(output_docx: Path) -> None:
    """Drop generated footer labels that Word would push onto blank extra pages."""
    document = Document(output_docx)
    for paragraph in list(document.paragraphs):
        label = " ".join(paragraph.text.split())
        if label.startswith("Sayfa ") and "/" in label:
            paragraph._element.getparent().remove(paragraph._element)
    for paragraph in reversed(list(document.paragraphs)):
        if paragraph.text.strip():
            break
        paragraph._element.getparent().remove(paragraph._element)
    document.save(output_docx)


def fix_multicolumn_footer_overflow(output_docx: Path) -> None:
    """Stop multi-column footer rows from pushing content onto spurious extra pages.

    pdf2docx renders a source page's footer row (address / page number / contact
    info, laid out side by side in the PDF) as a sequence of narrow DOCX sections
    that fake the horizontal layout using NEW_COLUMN section breaks. Word treats
    those breaks as hard page advances rather than same-page column moves, so a
    single-row footer can spawn its own near-blank page. Converting the break type
    to CONTINUOUS lets Word try to keep the row on the current page and only carry
    over genuine overflow, instead of always forcing a new page.
    """
    document = Document(output_docx)
    changed = False
    for section in document.sections:
        if section.start_type == WD_SECTION_START.NEW_COLUMN:
            section.start_type = WD_SECTION_START.CONTINUOUS
            changed = True
    if changed:
        document.save(output_docx)


def _section_list(body):
    """Group a document body's children into OOXML sections.

    Each entry is the list of body children belonging to that section, plus
    the section's continuation type and column count, read from the
    paragraph-level sectPr that closes it (or the body-level sectPr for the
    final section).
    """
    sections = []
    cur_start = 0
    children = list(body)
    for idx, child in enumerate(children):
        sect_pr = None
        if child.tag == qn("w:p"):
            p_pr = child.find(qn("w:pPr"))
            if p_pr is not None:
                sect_pr = p_pr.find(qn("w:sectPr"))
        elif child.tag == qn("w:sectPr"):
            sect_pr = child
        if sect_pr is None:
            continue
        end = idx + 1 if child.tag == qn("w:p") else idx
        cols_el = sect_pr.find(qn("w:cols"))
        num = int(cols_el.get(qn("w:num")) or "1") if cols_el is not None else 1
        type_el = sect_pr.find(qn("w:type"))
        ttype = type_el.get(qn("w:val")) if type_el is not None else "nextPage"
        sections.append({"elements": children[cur_start:end], "num": num, "ttype": ttype})
        cur_start = end
    return sections


def fix_footer_row_image_overflow(output_docx: Path, scale: float = 0.55) -> None:
    """Stop a footer's contact-info image from spilling onto its own blank page.

    pdf2docx renders a footer "row" that mixes short text (e.g. a contact name)
    with a small inline image (e.g. a QR code) as two separate, stacked
    paragraphs inside one continuous multi-column section. Word gives the text
    paragraph its own line height and then the image paragraph another line
    height on top of that, and near a page's bottom margin that combined
    height often doesn't fit even though the image alone would - so the image
    paragraph alone gets pushed onto its own near-blank page.

    Fix: merge the leading text paragraph into the image paragraph so they
    share a single line (the paragraph's height is then governed by the image
    alone, not text-height-plus-image-height), and shrink the image somewhat
    so that single line reliably fits in the tight space typically left at a
    page's bottom margin.
    """
    document = Document(output_docx)
    body = document.element.body
    sections = _section_list(body)

    changed = False
    for section in sections:
        if section["ttype"] != "continuous" or section["num"] < 2:
            continue
        paragraphs = [el for el in section["elements"] if el.tag == qn("w:p")]
        for p_text, p_img in zip(paragraphs, paragraphs[1:]):
            has_text_image = p_text.find(".//" + qn("wp:inline")) is not None
            has_img_image = p_img.find(".//" + qn("wp:inline")) is not None
            if has_text_image or not has_img_image:
                continue
            text_runs = [r for r in list(p_text) if r.tag == qn("w:r")]
            if not text_runs:
                continue
            for run in text_runs:
                for br in run.findall(qn("w:br")):
                    run.remove(br)

            p_pr_img = p_img.find(qn("w:pPr"))
            insert_at = list(p_img).index(p_pr_img) + 1 if p_pr_img is not None else 0
            for offset, run in enumerate(text_runs):
                p_img.insert(insert_at + offset, copy.deepcopy(run))

            img_runs = [r for r in list(p_img) if r.tag == qn("w:r")][len(text_runs):]
            if img_runs:
                space_run = copy.deepcopy(text_runs[-1])
                for br in space_run.findall(qn("w:br")):
                    space_run.remove(br)
                t_elements = space_run.findall(qn("w:t"))
                if t_elements:
                    t_elements[0].text = " "
                    t_elements[0].set(qn("xml:space"), "preserve")
                    for extra in t_elements[1:]:
                        space_run.remove(extra)
                    p_img.insert(insert_at + len(text_runs), space_run)

            body.remove(p_text)
            changed = True

    for section in sections:
        if section["ttype"] != "continuous" or section["num"] < 2:
            continue
        for element in section["elements"]:
            if element.getparent() is None:
                continue
            for inline in element.iter(qn("wp:inline")):
                extent = inline.find(qn("wp:extent"))
                if extent is None:
                    continue
                cy = int(extent.get("cy"))
                if cy < 150000:
                    continue
                cx = int(extent.get("cx"))
                new_cx, new_cy = int(cx * scale), int(cy * scale)
                extent.set("cx", str(new_cx))
                extent.set("cy", str(new_cy))
                for xfrm_ext in inline.iter(qn("a:ext")):
                    xfrm_ext.set("cx", str(new_cx))
                    xfrm_ext.set("cy", str(new_cy))
                changed = True

    if changed:
        document.save(output_docx)


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert a selectable-text PDF to editable DOCX locally.")
    parser.add_argument("input_pdf", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()

    original_bytes = args.input_pdf.read_bytes()
    pdf_bytes = original_bytes
    if pdf_bytes[:5] != PDF_HEADER:
        pdf_bytes = unwrap_signed_pdf(pdf_bytes)
        if pdf_bytes[:5] != PDF_HEADER:
            raise ValueError("The selected file is not a PDF.")

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="convertory-pdf2docx-") as tmp_dir:
        source_pdf = args.input_pdf
        if pdf_bytes is not original_bytes:
            source_pdf = Path(tmp_dir) / "input.pdf"
            source_pdf.write_bytes(pdf_bytes)

        converter = Converter(str(source_pdf))
        try:
            converter.convert(str(args.output_docx), multi_processing=False)
        finally:
            converter.close()
    remove_overflowing_page_labels(args.output_docx)
    fix_multicolumn_footer_overflow(args.output_docx)
    fix_footer_row_image_overflow(args.output_docx)
    if not args.output_docx.exists() or args.output_docx.stat().st_size < 1024:
        raise RuntimeError("The document engine did not produce a valid DOCX file.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print("Convertory DOCX conversion failed: " + str(error), file=sys.stderr)
        raise SystemExit(1)
